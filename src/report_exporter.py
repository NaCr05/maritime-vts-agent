"""
Export Markdown-style VTS reports to Word and PDF bytes.

The exporter intentionally supports a compact Markdown subset used by the app:
headings, bullet lists, paragraphs, and pipe tables.
"""
from __future__ import annotations

from io import BytesIO
from xml.sax.saxutils import escape


def markdown_to_docx_bytes(markdown_text: str) -> bytes:
    """Convert report Markdown to a DOCX file in memory."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx，请先安装依赖：pip install python-docx") from exc

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size in [("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
        styles[style_name].font.name = "Microsoft YaHei"
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.bold = True
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        if _is_table_start(lines, i):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _add_docx_table(doc, table_lines)
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(_strip_markdown_emphasis(line))
        i += 1

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def markdown_to_pdf_bytes(markdown_text: str) -> bytes:
    """Convert report Markdown to a PDF file in memory."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    except ImportError as exc:
        raise RuntimeError("缺少 reportlab，请先安装依赖：pip install reportlab") from exc

    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        base_font = "STSong-Light"
    except Exception:
        base_font = "Helvetica"

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "VTSBody",
        parent=styles["BodyText"],
        fontName=base_font,
        fontSize=9.5,
        leading=14,
        spaceAfter=6,
    )
    h1 = ParagraphStyle(
        "VTSH1",
        parent=body,
        fontSize=18,
        leading=24,
        alignment=1,
        textColor=colors.HexColor("#1f4e79"),
        spaceAfter=14,
    )
    h2 = ParagraphStyle(
        "VTSH2",
        parent=body,
        fontSize=13,
        leading=18,
        textColor=colors.HexColor("#1f4e79"),
        spaceBefore=8,
        spaceAfter=8,
    )
    h3 = ParagraphStyle(
        "VTSH3",
        parent=body,
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#1f4e79"),
        spaceBefore=6,
        spaceAfter=6,
    )

    story = []
    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            story.append(Spacer(1, 5))
            i += 1
            continue

        if _is_table_start(lines, i):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            table_data = _parse_markdown_table(table_lines)
            if table_data:
                story.append(_build_pdf_table(table_data, Table, TableStyle, colors, Paragraph, body))
                story.append(Spacer(1, 8))
            continue

        if line.startswith("# "):
            story.append(Paragraph(_xml_text(line[2:].strip()), h1))
        elif line.startswith("## "):
            story.append(Paragraph(_xml_text(line[3:].strip()), h2))
        elif line.startswith("### "):
            story.append(Paragraph(_xml_text(line[4:].strip()), h3))
        elif line.startswith("- "):
            story.append(Paragraph("• " + _xml_text(line[2:].strip()), body))
        else:
            story.append(Paragraph(_xml_text(_strip_markdown_emphasis(line)), body))
        i += 1

    bio = BytesIO()
    doc = SimpleDocTemplate(
        bio,
        pagesize=A4,
        leftMargin=1.4 * cm,
        rightMargin=1.4 * cm,
        topMargin=1.3 * cm,
        bottomMargin=1.3 * cm,
        title="VTS Monthly AIS Report",
    )
    doc.build(story)
    return bio.getvalue()


def _is_table_start(lines: list[str], index: int) -> bool:
    return (
        index + 1 < len(lines)
        and lines[index].strip().startswith("|")
        and lines[index + 1].strip().startswith("|")
        and set(lines[index + 1].strip().replace("|", "").replace(" ", "")) <= {"-", ":"}
    )


def _parse_markdown_table(table_lines: list[str]) -> list[list[str]]:
    if len(table_lines) < 2:
        return []
    rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if cells and all(set(c.replace(" ", "")) <= {"-", ":"} for c in cells):
            continue
        rows.append([_strip_markdown_emphasis(c) for c in cells])
    return rows


def _add_docx_table(doc, table_lines: list[str]) -> None:
    rows = _parse_markdown_table(table_lines)
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx].cells
        for c_idx, value in enumerate(row):
            cells[c_idx].text = value
            if r_idx == 0:
                for p in cells[c_idx].paragraphs:
                    for run in p.runs:
                        run.bold = True


def _build_pdf_table(table_data, Table, TableStyle, colors, Paragraph, body):
    pdf_data = [[Paragraph(_xml_text(cell), body) for cell in row] for row in table_data]
    table = Table(pdf_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9eaf7")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9aa6b2")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return table


def _strip_markdown_emphasis(text: str) -> str:
    return text.replace("**", "").replace("`", "")


def _xml_text(text: str) -> str:
    return escape(_strip_markdown_emphasis(text)).replace("\n", "<br/>")
