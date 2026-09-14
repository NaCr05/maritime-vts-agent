"""Exercise navigation and real monthly export callbacks through Streamlit."""
from io import BytesIO
from pathlib import Path
import os

import pandas as pd
import pytest
from docx import Document
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

SCRIPT = Path(__file__).with_name('smoke_app.py')


def start_page(section=None, direct=False):
    app = AppTest.from_file(str(SCRIPT), default_timeout=45)
    if section:
        app.session_state['main_nav'] = 'AIS 轨迹分析'
        app.session_state['ais_analysis_nav'] = section
    app.session_state['test_report_direct'] = direct
    app.run()
    assert not app.exception, [e.message for e in app.exception]
    assert not app.error, [e.value for e in app.error]
    return app


def test_rag_navigation_and_no_key_answer():
    app = start_page()
    assert any(h.value == '海事知识库问答' for h in app.subheader)
    assert len(app.chat_input) == 1
    app.chat_input[0].set_value('什么是 VTS？').run()
    assert not app.exception
    messages = app.session_state['messages']
    assert [m['role'] for m in messages] == ['user', 'assistant']
    assert messages[-1]['content']
    assert any('知识库检索路径' in item.value for item in app.success)


@pytest.mark.parametrize('section,widget_key', [
    ('📄 月度报告', 'monthly_report_month'),
    ('🛡️ 地理围栏', 'fence_zone_select'),
])
def test_analysis_navigation(section, widget_key):
    app = start_page(section)
    if section == '📄 月度报告':
        assert app.selectbox(widget_key).options
        assert app.button('generate_monthly_report')
    else:
        assert any('围栏' in item.value for item in app.markdown)
        assert len(app.multiselect) > 0


@pytest.mark.parametrize('direct', [True, False], ids=['injected-callbacks', 'full-navigation'])
def test_monthly_generation_and_all_downloads(direct):
    app = start_page('📄 月度报告', direct=direct)
    # Use the largest real month, and verify that other months are excluded.
    if direct:
        from src.ais_query import AISQuery
        data = AISQuery().df
    else:
        data = app.session_state['ais_data']
    month = data['time'].dt.strftime('%Y-%m').value_counts().index[0]
    app.selectbox('monthly_report_month').select(month).run()
    app.button('generate_monthly_report').click().run()
    assert not app.exception, [e.message for e in app.exception]
    assert not app.error, [e.value for e in app.error]
    payload = app.session_state['monthly_report_payload']
    expected = data[data['time'].dt.strftime('%Y-%m') == month]
    pd.testing.assert_frame_equal(payload['month_df'], expected)
    assert payload['monthly_vessels']['轨迹点数'].sum() == len(expected)
    assert payload['monthly_vessels']['船 ID'].nunique() == expected['vessel_id'].nunique()
    assert payload['anomalies'] is not None
    assert payload['violations'] is not None
    downloads = app.session_state['test_downloads']
    assert set(downloads) == {
        'download_monthly_report_md', 'download_monthly_report_docx', 'download_monthly_report_pdf',
    }
    assert downloads['download_monthly_report_md'].decode('utf-8-sig') == payload['text']
    doc = Document(BytesIO(downloads['download_monthly_report_docx']))
    assert month in doc.paragraphs[0].text
    assert len(doc.tables) >= 4
    pdf = PdfReader(BytesIO(downloads['download_monthly_report_pdf']))
    assert len(pdf.pages) >= 1
    assert month in ''.join(page.extract_text() for page in pdf.pages)
    # A rerun retains the generated report and its downloads.
    app.run()
    assert not app.exception and app.session_state['monthly_report_payload']['text'] == payload['text']
    output = os.environ.get('VTS_TEST_OUTPUT_DIR')
    if output and not direct:
        folder = Path(output)
        folder.mkdir(parents=True, exist_ok=True)
        for suffix in ['md', 'docx', 'pdf']:
            (folder / f'monthly-report.{suffix}').write_bytes(downloads[f'download_monthly_report_{suffix}'])


def test_no_key_agent_explains_configuration_without_fake_fallback():
    app = start_page('自然语言查询')
    warnings = '\n'.join(item.value for item in app.warning)
    assert 'LLM_API_KEY' in warnings
    assert '已使用关键词' not in warnings
    app.button('nl_ex_2').click().run()
    assert not app.exception
    text = '\n'.join(item.value for item in app.markdown)
    assert 'LLM_API_KEY' in text
    assert 'DEEPSEEK_API_KEY' not in text
